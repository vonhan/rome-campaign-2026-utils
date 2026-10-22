from __future__ import annotations

import re
from zipfile import BadZipFile
from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from src.diagnostics import ValidationIssue, WorkbookValidationError
from src.models import Appointment, DayAssignment, Volunteer


_DATE_RE = re.compile(r"^\d{1,2}/\d{1,2}(?:\s|$)")
_NEW_VOLUNTEER_RE = re.compile(
    r"^(?P<full_name>[^\d,;:]+?\S)\s+(?P<phone>\d+(?:\s+\d+)*)$",
)
_UNION_RE = re.compile(
    r"^(?:"
    r"(?P<direct>[A-Z0-9]+)"
    r"|\((?:include|aggregato)\s+(?P<wrapped>[A-Z0-9]+)\)"
    r")$",
    flags=re.IGNORECASE,
)
_FIELD_ALIASES = {
    "group_code": ("codice gruppo", "codice", "gruppo", "code"),
    "group_contact": ("referente gruppo", "referente", "group contact"),
    "group_contact_phone": (
        "telefono referente gruppo",
        "telefono referente",
        "telefono capogruppo",
        "telefono gruppo",
        "phone referente",
        "ref phone",
        "referente telefono",
        "group contact phone",
    ),
    "group_identification_code": (
        "cod id gruppo",
        "cod identificativo gruppo",
        "codice identificativo gruppo",
        "codice identificativo per riconoscimento",
    ),
    "kingdom_hall": ("sala del regno", "sala", "regno", "kingdom hall"),
    "guest_count": ("n ospiti", "numero ospiti", "ospiti", "guests"),
    "guest_vehicle": ("auto ospiti", "auto", "veicolo ospiti", "car"),
}


def _clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_header(value: object) -> str:
    text = _clean(value).lower().replace("_", " ")
    text = re.sub(r"[^\w\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def parse_volunteer(value: object) -> Volunteer | None:
    raw = re.sub(r"\s+", " ", _clean(value))
    if not raw:
        return None

    match = _NEW_VOLUNTEER_RE.fullmatch(raw)
    if not match:
        raise ValueError(
            "Formato volontario non valido; usare NOME COMPLETO TELEFONO, "
            "con telefono finale composto da sole cifre."
        )
    full_name = match.group("full_name").strip()
    phone = re.sub(r"\s+", "", match.group("phone"))
    return Volunteer(
        raw=f"{full_name} {phone}",
        full_name=full_name,
        phone=phone,
    )


def parse_volunteers(value: object) -> tuple[Volunteer, ...]:
    raw = _clean(value)
    if not raw:
        return ()

    parts = [part.strip() for part in raw.split(";")]
    if parts and not parts[-1]:
        parts.pop()
    if any(not part for part in parts):
        raise ValueError(
            "Separatore ';' non valido: ogni separatore deve trovarsi tra due volontari."
        )

    volunteers: list[Volunteer] = []
    seen: set[tuple[str, str]] = set()
    for part in parts:
        volunteer = parse_volunteer(part)
        assert volunteer is not None
        key = volunteer.identity_key
        if key in seen:
            raise ValueError("Lo stesso volontario è ripetuto nella medesima cella.")
        seen.add(key)
        volunteers.append(volunteer)
    return tuple(volunteers)


def parse_union_group(value: object) -> tuple[str, ...]:
    raw = re.sub(r"\s+", " ", _clean(value))
    if not raw:
        return ()

    codes: list[str] = []
    for item in raw.split(";"):
        item = item.strip()
        match = _UNION_RE.fullmatch(item)
        if not match:
            raise ValueError(
                "Formato UNIONI non valido; usare uno o più codici separati "
                "da ';', con CODICE, (include CODICE) o (aggregato CODICE)."
            )
        code = (match.group("direct") or match.group("wrapped")).upper()
        if code in codes:
            raise ValueError("Il codice UNIONI è ripetuto nella medesima cella.")
        codes.append(code)
    return tuple(codes)


def _find_date_columns(header_row: Sequence[object]) -> list[int]:
    return [
        index
        for index, value in enumerate(header_row)
        if _DATE_RE.match(_clean(value))
    ]


def _find_base_columns(
    header_row: Sequence[object],
    first_date_column: int,
) -> dict[str, int]:
    normalized = [_normalize_header(value) for value in header_row]
    result: dict[str, int] = {}
    missing: list[str] = []
    required_fields = {
        "group_code",
        "group_contact",
        "group_identification_code",
        "kingdom_hall",
        "guest_count",
        "guest_vehicle",
    }
    for field, aliases in _FIELD_ALIASES.items():
        index = next(
            (position for position, value in enumerate(normalized) if value in aliases),
            None,
        )
        if index is None:
            continue
        result[field] = index

    if first_date_column != 7:
        raise ValueError(
            "Formato Excel non valido: è richiesto lo schema v3 con "
            "'Cod. ID. Gruppo' in colonna D e il primo blocco data in colonna H."
        )

    definitive_positions = {
        "group_code": 0,
        "group_contact": 1,
        "group_contact_phone": 2,
        "group_identification_code": 3,
        "kingdom_hall": 4,
        "guest_count": 5,
        "guest_vehicle": 6,
    }
    recognized_positions_match = all(
        result.get(field, position) == position
        for field, position in definitive_positions.items()
    )
    if recognized_positions_match:
        result.setdefault("group_contact_phone", 2)
        result.setdefault("group_identification_code", 3)
        result.setdefault("kingdom_hall", 4)

    for field in required_fields:
        if field not in result:
            missing.append(_FIELD_ALIASES[field][0])
    if missing:
        raise ValueError(
            "Intestazioni obbligatorie mancanti: " + ", ".join(missing) + "."
        )
    return result


def _find_assignment_columns(
    second_row: Sequence[object], start: int, end: int, date_label: str
) -> tuple[int, int, int, int | None]:
    block = {
        _normalize_header(second_row[index]): index
        for index in range(start, min(end, len(second_row)))
        if _normalize_header(second_row[index])
    }
    morning = block.get("mattina")
    afternoon = block.get("pomeriggio")
    union = block.get("unioni")
    note = block.get("note")
    if morning is None or afternoon is None or union is None:
        raise ValueError(
            f"Il blocco del {date_label} deve contenere "
            "Mattina, Pomeriggio e UNIONI."
        )
    return morning, afternoon, union, note


def read_appointments(path: str | Path) -> list[Appointment]:
    workbook_path = Path(path)
    if workbook_path.suffix.lower() != ".xlsx":
        raise ValueError("Selezionare un file Excel con estensione .xlsx.")

    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except (OSError, ValueError, BadZipFile, InvalidFileException) as exc:
        raise ValueError(f"Impossibile leggere il file Excel: {exc}") from exc

    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        if len(rows) < 2:
            raise ValueError("Il file Excel non contiene le due righe di intestazione.")

        header_row = rows[0]
        second_row = rows[1]
        date_columns = _find_date_columns(header_row)
        if not date_columns:
            raise ValueError(
                "Nessun blocco data trovato nella prima riga di intestazione."
            )
        base_columns = _find_base_columns(header_row, date_columns[0])

        assignment_columns: list[tuple[str, int, int, int, int | None]] = []
        for position, date_column in enumerate(date_columns):
            block_end = (
                date_columns[position + 1]
                if position + 1 < len(date_columns)
                else len(header_row)
            )
            date_label = _clean(header_row[date_column])
            columns = _find_assignment_columns(
                second_row, date_column, block_end, date_label
            )
            assignment_columns.append((date_label, *columns))

        appointments: list[Appointment] = []
        issues: list[ValidationIssue] = []
        for row_number, row in enumerate(rows[2:], start=3):
            base_values = {
                field: _clean(row[index]) if index < len(row) else ""
                for field, index in base_columns.items()
            }
            if not any(base_values.values()):
                continue

            assignments: list[DayAssignment] = []
            for date_label, morning, afternoon, union, note in assignment_columns:
                parsed_slots: dict[str, tuple[Volunteer, ...]] = {}
                for field, column in (
                    ("Mattina", morning),
                    ("Pomeriggio", afternoon),
                ):
                    try:
                        parsed_slots[field] = parse_volunteers(
                            row[column] if column < len(row) else None
                        )
                    except ValueError as exc:
                        issues.append(
                            ValidationIssue(
                                sheet=worksheet.title,
                                cell=f"{get_column_letter(column + 1)}{row_number}",
                                row_number=row_number,
                                date_label=date_label,
                                field=field,
                                reason=str(exc),
                            )
                        )
                        parsed_slots[field] = ()
                try:
                    union_group_code = parse_union_group(
                        row[union] if union < len(row) else None
                    )
                except ValueError as exc:
                    issues.append(
                        ValidationIssue(
                            sheet=worksheet.title,
                            cell=f"{get_column_letter(union + 1)}{row_number}",
                            row_number=row_number,
                            date_label=date_label,
                            field="UNIONI",
                            reason=str(exc),
                        )
                    )
                    union_group_code = ()
                assignments.append(
                    DayAssignment(
                        date_label=date_label,
                        morning=parsed_slots["Mattina"],
                        afternoon=parsed_slots["Pomeriggio"],
                        union_group_code=union_group_code,
                        note=(
                            _clean(row[note])
                            if note is not None and note < len(row)
                            else ""
                        ),
                    )
                )
            appointments.append(
                Appointment(
                    **base_values,
                    assignments=assignments,
                    source_row_number=row_number,
                )
            )

        known_group_codes = {
            appointment.group_code.upper()
            for appointment in appointments
            if appointment.group_code
        }
        for appointment in appointments:
            for assignment, assignment_columns_for_date in zip(
                appointment.assignments,
                assignment_columns,
                strict=True,
            ):
                date_label, _morning, _afternoon, union, _note = (
                    assignment_columns_for_date
                )
                for code in assignment.union_group_code:
                    if code.upper() not in known_group_codes:
                        issues.append(
                            ValidationIssue(
                                sheet=worksheet.title,
                                cell=(
                                    f"{get_column_letter(union + 1)}"
                                    f"{appointment.source_row_number}"
                                ),
                                row_number=appointment.source_row_number,
                                date_label=date_label,
                                field="UNIONI",
                                reason=(
                                    f"Il codice gruppo '{code}' non è presente "
                                    "nella colonna Codice Gruppo."
                                ),
                            )
                        )
        if issues:
            raise WorkbookValidationError(issues)
        return appointments
    finally:
        workbook.close()
