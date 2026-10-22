from __future__ import annotations

from collections.abc import Callable
from pathlib import Path

import pytest
from openpyxl import Workbook


def _expand_compact_row(row: list[object]) -> list[object]:
    """Convert the compact test row into the v3 workbook layout."""
    if not any(value is not None and value != "" for value in row):
        return row
    if len(row) == 11:
        return [
            row[0],
            row[1],
            "",
            f"{row[0]}-ID",
            row[2],
            row[3],
            row[4],
            row[5],
            row[6],
            "",
            row[7],
            row[8],
            row[9],
            "",
            row[10],
        ]
    return row


@pytest.fixture
def synthetic_workbook() -> Callable[[Path, list[list[object]] | None], Path]:
    def create(
        path: Path,
        data_rows: list[list[object]] | None = None,
    ) -> Path:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Appuntamenti"
        worksheet.append(
            [
                "Codice Gruppo",
                "Referente Gruppo",
                "Telefono Referente",
                "Cod. ID. Gruppo",
                "Sala del Regno",
                "N. Ospiti",
                "Auto Ospiti",
                "26/10 (Lunedì)",
                "",
                "",
                "",
                "27/10 (Martedì)",
                "",
                "",
                "",
            ]
        )
        worksheet.append(
            [
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "Mattina",
                "Pomeriggio",
                "UNIONI",
                "NOTE",
                "Mattina",
                "Pomeriggio",
                "UNIONI",
                "NOTE",
            ]
        )
        worksheet.merge_cells("H1:K1")
        worksheet.merge_cells("L1:O1")
        rows = data_rows or [
            [
                "RM-01",
                "Referente Sintetico",
                "Sala Centro",
                4,
                "Auto blu",
                "ROSSI MARIO 000111222",
                "VERDI ANNA 000333444",
                "Portare il materiale",
                "BIANCHI LUCA 000555666",
                "",
                "Ingresso laterale",
            ]
        ]
        for row in rows:
            worksheet.append(_expand_compact_row(row))
        workbook.save(path)
        return path

    return create


@pytest.fixture
def synthetic_workbook_v2() -> Callable[[Path, list[list[object]] | None], Path]:
    def create(
        path: Path,
        data_rows: list[list[object]] | None = None,
    ) -> Path:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Appuntamenti"
        worksheet.append(
            [
                "Codice Gruppo",
                "Referente Gruppo",
                "Telefono Referente",
                "Cod. ID. Gruppo",
                "Sala del Regno",
                "N. Ospiti",
                "Auto Ospiti",
                "26/10 (Lunedì)",
                "",
                "",
                "27/10 (Martedì)",
                "",
                "",
            ]
        )
        worksheet.append(
            [
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "Mattina",
                "Pomeriggio",
                "UNIONI",
                "Mattina",
                "Pomeriggio",
                "UNIONI",
            ]
        )
        worksheet.merge_cells("H1:J1")
        worksheet.merge_cells("K1:M1")
        rows = data_rows or [
            [
                "RM-01",
                "Referente Sintetico",
                "+39 333 444 5555",
                "Sala Centro",
                4,
                "Auto blu",
                "ROSSI MARIO 0001112222; VERDI PAOLO 0003334444",
                "",
                "",
                "BIANCHI LUCA 0005556666",
                "",
                "",
            ]
        ]
        for row in rows:
            worksheet.append(
                [
                    row[0],
                    row[1],
                    row[2],
                    f"{row[0]}-ID",
                    *row[3:],
                ]
            )
        workbook.save(path)
        return path

    return create
