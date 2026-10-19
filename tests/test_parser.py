from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from src.diagnostics import WorkbookValidationError
from src.excel_parser import (
    parse_union_group,
    parse_volunteer,
    parse_volunteers,
    read_appointments,
)


def test_read_appointments_understands_definitive_header_shape(
    tmp_path: Path, synthetic_workbook_v2
) -> None:
    workbook_path = synthetic_workbook_v2(tmp_path / "synthetic.xlsx")

    appointments = read_appointments(workbook_path)

    assert len(appointments) == 1
    appointment = appointments[0]
    assert appointment.group_code == "RM-01"
    assert appointment.group_contact == "Referente Sintetico"
    assert appointment.group_contact_phone == "+39 333 444 5555"
    assert appointment.kingdom_hall == "Sala Centro"
    assert appointment.guest_count == "4"
    assert appointment.guest_vehicle == "Auto blu"
    assert appointment.source_row_number == 3
    assert [item.date_label for item in appointment.assignments] == [
        "26/10 (Lunedì)",
        "27/10 (Martedì)",
    ]
    assert appointment.assignments[0].union_group_code == "JOINED001"
    assert [
        (volunteer.full_name, volunteer.phone)
        for volunteer in appointment.assignments[0].morning
    ] == [
        ("ROSSI MARIO", "0001112222"),
        ("VERDI PAOLO", "0003334444"),
    ]
    assert appointment.assignments[1].afternoon == ()


def test_definitive_base_columns_use_positions_for_unknown_c_and_d_headers(
    tmp_path: Path, synthetic_workbook_v2
) -> None:
    path = synthetic_workbook_v2(tmp_path / "positional.xlsx")

    workbook = load_workbook(path)
    sheet = workbook.active
    sheet["C1"] = "Intestazione telefono personalizzata"
    sheet["D1"] = "Intestazione luogo personalizzata"
    workbook.save(path)
    workbook.close()

    appointment = read_appointments(path)[0]

    assert appointment.group_contact_phone == "+39 333 444 5555"
    assert appointment.kingdom_hall == "Sala Centro"


def test_read_appointments_keeps_every_non_empty_data_row(
    tmp_path: Path, synthetic_workbook
) -> None:
    rows = [
        ["DUP", "Primo", "Sala A", 1, "", "", "", "", "", "", ""],
        [None] * 11,
        ["DUP", "Secondo", "Sala B", 2, "", "", "", "", "", "", ""],
    ]
    workbook_path = synthetic_workbook(tmp_path / "rows.xlsx", rows)

    appointments = read_appointments(workbook_path)

    assert [item.source_row_number for item in appointments] == [3, 5]
    assert len(appointments) == 2
    assert all(len(item.assignments) == 2 for item in appointments)


@pytest.mark.parametrize(
    ("raw", "name", "phone"),
    [
        ("ROSSI MARIO 000123456", "ROSSI MARIO", "000123456"),
        (
            "DE ROSSI MARIA LUISA 0009990000",
            "DE ROSSI MARIA LUISA",
            "0009990000",
        ),
    ],
)
def test_parse_volunteer_preserves_full_name(
    raw: str, name: str, phone: str
) -> None:
    volunteer = parse_volunteer(raw)

    assert volunteer is not None
    assert volunteer.full_name == name
    assert volunteer.phone == phone
    assert volunteer.display == raw


def test_parse_volunteer_returns_none_for_empty_value() -> None:
    assert parse_volunteer("  ") is None


def test_parse_volunteers_splits_only_on_semicolon() -> None:
    volunteers = parse_volunteers(
        "ROSSI MARIO 0000000001; VERDI PAOLO 0000000002"
    )

    assert [(item.full_name, item.phone) for item in volunteers] == [
        ("ROSSI MARIO", "0000000001"),
        ("VERDI PAOLO", "0000000002"),
    ]


@pytest.mark.parametrize(
    "raw",
    [
        "ROSSI, MARIO: 0000000001",
        "ROSSI MARIO 0000000001 e VERDI PAOLO 0000000002",
        "ROSSI MARIO +39 0000000001",
        "ROSSI MARIO abc",
        "0000000001",
    ],
)
def test_invalid_volunteer_formats_are_rejected(raw: str) -> None:
    with pytest.raises(ValueError, match="Formato volontario non valido"):
        parse_volunteers(raw)


@pytest.mark.parametrize(
    "raw",
    [
        "; ROSSI MARIO 0000000001",
        "ROSSI MARIO 0000000001;; VERDI PAOLO 0000000002",
        "ROSSI MARIO 0000000001;;",
    ],
)
def test_empty_volunteer_elements_are_rejected(raw: str) -> None:
    with pytest.raises(ValueError, match="Separatore ';' non valido"):
        parse_volunteers(raw)


def test_single_trailing_semicolon_is_accepted() -> None:
    volunteers = parse_volunteers("ROSSI MARIO 0000000001;")

    assert [(item.full_name, item.phone) for item in volunteers] == [
        ("ROSSI MARIO", "0000000001")
    ]


def test_duplicate_volunteer_in_same_cell_is_rejected() -> None:
    with pytest.raises(ValueError, match="ripetuto"):
        parse_volunteers(
            "ROSSI MARIO 0000000001;  rossi   mario 0000000001"
        )


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("UNITO0000001", "UNITO0000001"),
        ("(include UNITO0000001)", "UNITO0000001"),
        ("(aggregato unito0000001)", "UNITO0000001"),
        ("", ""),
        (None, ""),
    ],
)
def test_parse_union_group_accepts_documented_forms(
    raw: object, expected: str
) -> None:
    assert parse_union_group(raw) == expected


@pytest.mark.parametrize(
    "raw",
    [
        "(include UNITO0000001",
        "include UNITO0000001",
        "(unito UNITO0000001)",
        "UNITO0000001; ALTRO000001",
        "(include UNITO0000001 ALTRO000001)",
    ],
)
def test_parse_union_group_rejects_ambiguous_values(raw: str) -> None:
    with pytest.raises(ValueError, match="Formato UNIONI non valido"):
        parse_union_group(raw)


def test_validation_errors_are_aggregated_with_coordinates_without_values(
    tmp_path: Path, synthetic_workbook_v2
) -> None:
    sensitive_name = "PERSONA SEGRETA"
    rows = [
        [
            "RM-ERR",
            "Referente Sintetico",
            "0000000000",
            "Sala",
            1,
            "",
            f"{sensitive_name} telefono",
            "ALTRO NOME 123;;",
            "(include BAD CODE)",
            "",
            "",
            "",
        ]
    ]
    path = synthetic_workbook_v2(tmp_path / "errors.xlsx", rows)

    with pytest.raises(WorkbookValidationError) as raised:
        read_appointments(path)

    error = raised.value
    assert len(error.issues) == 3
    assert [issue.cell for issue in error.issues] == ["G3", "H3", "I3"]
    message = str(error)
    assert "Foglio 'Appuntamenti'" in message
    assert "26/10 (Lunedì)" in message
    assert sensitive_name not in message


def test_missing_required_header_is_reported(tmp_path: Path) -> None:
    path = tmp_path / "invalid.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        ["Codice Gruppo", "Referente Gruppo", "Sala del Regno", "26/10"]
    )
    sheet.append(["", "", "", "Mattina"])
    workbook.save(path)

    with pytest.raises(ValueError, match="Intestazioni obbligatorie mancanti"):
        read_appointments(path)


def test_incomplete_date_block_is_reported(tmp_path: Path) -> None:
    path = tmp_path / "invalid-block.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "Codice Gruppo",
            "Referente Gruppo",
            "Sala del Regno",
            "N. Ospiti",
            "Auto Ospiti",
            "26/10 (Lunedì)",
            "",
        ]
    )
    sheet.append(["", "", "", "", "", "Mattina", "Pomeriggio"])
    workbook.save(path)

    with pytest.raises(
        ValueError,
        match="Mattina, Pomeriggio e UNIONI",
    ):
        read_appointments(path)


def test_non_xlsx_file_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "appointments.csv"
    path.write_text("synthetic", encoding="utf-8")

    with pytest.raises(ValueError, match="estensione .xlsx"):
        read_appointments(path)


def test_corrupted_xlsx_is_reported_clearly(tmp_path: Path) -> None:
    path = tmp_path / "corrupted.xlsx"
    path.write_text("not an Excel archive", encoding="utf-8")

    with pytest.raises(ValueError, match="Impossibile leggere il file Excel"):
        read_appointments(path)
